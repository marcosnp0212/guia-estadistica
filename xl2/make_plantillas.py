import sys; sys.path.insert(0, '/home/claude/xl2')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import FormulaRule
from engine import F_DATA, F_IN, F_HEAD, BOX, INK, TEMA

F_OUT = PatternFill('solid', fgColor='E2F2EA')
wb = Workbook(); wb.remove(wb.active)
TOC = []


def sheet(name, title, color, desc, how):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = TEMA[color]
    ws.sheet_view.zoomScale = 115
    for col in range(1, 9):
        ws.cell(1, col).fill = PatternFill('solid', fgColor=TEMA[color])
    c = ws.cell(1, 1, title); c.font = Font(bold=True, size=15, color='FFFFFF'); ws.row_dimensions[1].height = 30
    c.alignment = Alignment(vertical='center')
    ws.cell(2, 1, desc).font = Font(italic=True, color='4F5D65')
    ws.cell(3, 1, how).font = Font(size=10, color='6B5800')
    ws.cell(4, 1, 'Amarillo: escribe o pega aquí tus datos.  Verde: resultados (no los toques).').font = Font(size=9, color='7A878E')
    TOC.append((name, title, desc))
    return ws


def head(ws, r, c, labels):
    for j, t in enumerate(labels):
        x = ws.cell(r, c + j, t); x.font = Font(bold=True, color=INK); x.fill = F_HEAD; x.border = BOX
        x.alignment = Alignment(horizontal='center', wrap_text=True)


def inp(ws, rng):
    for row in ws[rng]:
        for c in row:
            c.fill = F_IN; c.border = BOX


def out(ws, r, c, label, formula, fmt='0.0000', note=None):
    a = ws.cell(r, c, label); a.font = Font(bold=True)
    if isinstance(formula, str) and formula.startswith('='):
        formula = '=IFERROR(' + formula[1:] + ',"")'
    b = ws.cell(r, c + 1, formula); b.fill = F_OUT; b.border = BOX; b.number_format = fmt; b.font = Font(bold=True)
    if note:
        ws.cell(r, c + 2, note).font = Font(italic=True, size=9, color='7A878E')
    return f'{CL(c + 1)}{r}'


def widths(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


# ======================================================= Traductor
ws = wb.create_sheet('Traductor')
ws.sheet_properties.tabColor = '1F2A30'
widths(ws, {'A': 46, 'B': 46, 'C': 40, 'D': 18})
ws['A1'] = '¿Qué me están pidiendo?  Traductor de enunciados'; ws['A1'].font = Font(bold=True, size=15, color=INK)
ws['A2'] = 'Busca en el enunciado una frase parecida (Ctrl + B para buscar) y mira qué hay que calcular y con qué.'; ws['A2'].font = Font(italic=True, color='4F5D65')
rows = [
    ('TEMA 1 · UNA VARIABLE', None, None, None),
    ('"valor medio", "promedio", "por término medio"', 'Media', '=PROMEDIO(datos) · tabla: SUMAPRODUCTO(x;n)/SUMA(n)', 'A_Discreta / C_Intervalos'),
    ('"el más frecuente", "el más habitual", "el que más se repite"', 'Moda (con intervalos de distinta amplitud: por densidad)', '=MODA.UNO(datos) · intervalos: fórmula de la moda', 'C_Intervalos'),
    ('"el valor central", "la mitad tiene menos de…", "deja el 50 % por debajo"', 'Mediana', '=MEDIANA(datos) · tabla: primera Nᵢ ≥ n/2', 'B_Tabla_valores'),
    ('"el 25 % que menos…", "a partir de cuánto se está en el 25 % inferior"', 'Primer cuartil (P25)', '=CUARTIL.INC(datos;1)', 'A / B / C'),
    ('"el 25 % que más…", "estar entre el 25 % que más gasta"', 'Tercer cuartil (P75)', '=CUARTIL.INC(datos;3)', 'A / B / C'),
    ('"el X % que más…", "estar en el X % superior"', 'Percentil 100 − X', '=PERCENTIL.INC(datos;(100−X)/100)', 'A / B / C'),
    ('"el X % que menos…", "como mucho para estar en el X % inferior"', 'Percentil X', '=PERCENTIL.INC(datos;X/100)', 'A / B / C'),
    ('"entre qué valores está el 50 % central"', 'C₁ y C₃', 'CUARTIL.INC 1 y 3', 'A / B / C'),
    ('"entre qué valores está el X % central"', 'P(50−X/2) y P(50+X/2)  (30 % central: P35 y P65)', 'PERCENTIL.INC', 'C_Intervalos'),
    ('"cuántos tienen más de k" / "al menos k" / "menos de k"', '"más de" NO incluye k; "al menos" SÍ lo incluye', '=CONTAR.SI(datos;">k") · ">=k" · "<k"', 'A_Discreta'),
    ('"dispersión", "variabilidad absoluta", "cuánto se separan"', 'Varianza o desviación típica', '=VAR.P · =DESVEST.P', 'A / B / C'),
    ('"más homogéneo", "más representativa la media", "variabilidad relativa", comparar grupos con medias distintas', 'Coeficiente de variación (el MENOR es el más homogéneo)', '=DESVEST.P/PROMEDIO', 'A / C / D_Grupos'),
    ('"cuasivarianza", "varianza muestral corregida"', 'Cuasivarianza (divide entre n − 1)', '=VAR.S · =DESVEST.S', 'A_Discreta'),
    ('"forma", "simétrica", "cola"', 'Coeficiente de asimetría (signo)', '=COEFICIENTE.ASIMETRIA.P', 'A_Discreta'),
    ('"picuda", "aplanada", "apuntamiento"', 'Curtosis (signo)', '=CURTOSIS', 'A_Discreta'),
    ('"valores atípicos", "outliers", "diagrama de caja"', 'Límites C₁ − 1,5·IQR y C₃ + 1,5·IQR', 'Fórmula con CUARTIL.INC', 'A_Discreta'),
    ('"comparar la posición de alguien en dos grupos distintos"', 'Tipificar: z = (x − media)/S', '=NORMALIZACION(x;media;S)', '—'),
    ('"media de todos" a partir de medias de grupos de distinto tamaño', 'Media ponderada', '=SUMAPRODUCTO(medias;tamaños)/SUMA(tamaños)', 'D_Grupos'),
    ('"varianza de toda la empresa" a partir de grupos', 'Varianza dentro + varianza entre', '—', 'D_Grupos'),
    ('"% de la variabilidad debido a las diferencias entre grupos", "desigualdad"', 'Varianza entre ÷ varianza total', '—', 'D_Grupos'),
    ('"todos suben X €" / "suben un X %" ¿qué pasa con media y varianza?', 'Sumar: media + X, varianza igual. Multiplicar por a: media·a, varianza·a²', '—', 'E_Transformacion'),
    ('"sin calcularlo, indica qué pasa…"', 'Reglas de cambio de origen y escala (arriba)', '—', 'E_Transformacion'),
    ('TEMA 2 · DOS VARIABLES', None, None, None),
    ('"distribución marginal de X"', 'Los totales de las filas (o de las columnas para Y)', '=SUMA(fila)', 'F_Doble_entrada'),
    ('"distribución de Y condicionada a X = …", "de los que…, ¿qué %…?"', 'Esa fila (o columna) dividida entre SU total', '=casilla/total de la fila', 'F_Doble_entrada'),
    ('"¿son independientes?", "¿influye X en Y?", "¿depende…?"', 'Comparar fᵢⱼ con fᵢ.·f.ⱼ en todas las casillas', '—', 'F_Doble_entrada'),
    ('"frecuencias relativas" de una tabla de doble entrada', 'Cada casilla ÷ n', '=casilla/$total$', 'F_Doble_entrada'),
    ('TEMA 3 · REGRESIÓN', None, None, None),
    ('"recta que explica Y a partir de X", "Y en función de X"', 'Recta de Y sobre X: Ŷ = a + b·X', '=PENDIENTE(Y;X) · =INTERSECCION.EJE(Y;X)', 'G_Regresion'),
    ('"cuánto cambia Y por cada unidad de X"', 'La pendiente b', '=PENDIENTE(Y;X)', 'G / H'),
    ('"grado de relación lineal", "intensidad y sentido"', 'Correlación r', '=COEF.DE.CORREL(X;Y)', 'G / H'),
    ('"bondad del ajuste", "% explicado", "fiabilidad de la recta"', 'R² (en %)', '=COEFICIENTE.R2(Y;X)', 'G / H'),
    ('"% NO explicado"', '1 − R²', '—', 'G / H'),
    ('"¿cuánto se espera de Y si X vale…?", "estima", "predice"', 'Sustituir X en la recta', '=PRONOSTICO.LINEAL(x;Y;X)', 'G / H'),
    ('"variación conjunta", "¿relación directa o inversa?"', 'Covarianza (signo)', '=COVARIANZA.P(X;Y)', 'G_Regresion'),
    ('Te dan medias, S y Sxy (o r) pero no los datos', 'b = Sxy/Sx² · a = ȳ − b·x̄ · r = Sxy/(Sx·Sy)', '—', 'H_Reg_resumen'),
    ('TEMA 4 · ÍNDICES', None, None, None),
    ('"índice con base 20XX", "evolución respecto a 20XX"', 'Índice simple: valor/valor base × 100', '=B3/$B$2*100', 'I_Indices'),
    ('"tasa de variación", "cuánto ha subido/bajado"', 'valor/anterior − 1', '=B3/B2-1', 'I_Indices'),
    ('"cambia la base a 20XX"', 'Dividir toda la serie entre el índice de 20XX', '=C2/$C$6*100', 'I_Indices'),
    ('"subió un X % y luego bajó un Y %" (acumulado)', 'Multiplicar: (1+X)·(1−Y)', '=PRODUCTO', 'J_Tasas'),
    ('"tasa media anual", "crecimiento medio por año"', 'Media geométrica: (1+T)^(1/años) − 1', '=MEDIA.GEOM · =(1+T)^(1/k)-1', 'I / J'),
    ('"tasa mensual equivalente" a una anual', '(1 + anual)^(1/12) − 1', '=(1+t)^(1/12)-1', 'J_Tasas'),
    ('"A superó en un X % a B"', 'A = (1 + X/100)·B → B = A/(1 + X/100)', '—', 'J_Tasas'),
    ('"índice de precios con las cantidades del año base"', 'Laspeyres', 'SUMAPRODUCTO(pt;q0)/SUMAPRODUCTO(p0;q0)*100', 'K_Laspeyres'),
    ('"con las cantidades actuales"', 'Paasche', 'SUMAPRODUCTO(pt;qt)/SUMAPRODUCTO(p0;qt)*100', 'K_Laspeyres'),
    ('"índice de Fisher"', '√(Laspeyres · Paasche)', '=RAIZ(L*P)', 'K_Laspeyres'),
    ('"en términos reales", "euros constantes", "poder adquisitivo", "deflactar"', 'Valor nominal ÷ IPC × 100', '=B2/C2*100', 'L_Deflactar'),
    ('"¿está bien actualizado el salario/alquiler con el IPC?"', 'Comparar con valor × IPC nuevo/IPC viejo', '—', 'L_Deflactar'),
]
head(ws, 4, 1, ['Si el enunciado dice…', 'Te están pidiendo…', 'En Excel', 'Plantilla'])
r = 5
for a, b, c, d in rows:
    if b is None:
        x = ws.cell(r, 1, a); x.font = Font(bold=True, color='FFFFFF')
        for j in range(1, 5):
            ws.cell(r, j).fill = PatternFill('solid', fgColor='5B6B78')
    else:
        for j, v in enumerate([a, b, c, d]):
            x = ws.cell(r, 1 + j, v); x.alignment = Alignment(wrap_text=True, vertical='top'); x.border = BOX
            if j == 2:
                x.font = Font(name='Consolas', size=9, color='1D6F42')
                if isinstance(v, str) and v.startswith('='):
                    x.data_type = 's'
    r += 1
ws.freeze_panes = 'A5'

N = 500  # filas de datos

# ======================================================= A_Discreta
ws = sheet('A_Discreta', 'A · Datos sueltos (una variable)', 't1', 'Pega tus datos en la columna A (desde A7). Calcula tabla de frecuencias y todas las medidas.',
           'Los valores de la tabla (columna D) salen solos si son números enteros seguidos (0, 1, 2…). Si no, bórralos y escribe tú los valores distintos.')
widths(ws, {'A': 12, 'B': 3, 'C': 3, 'D': 10, 'E': 10, 'F': 10, 'G': 10, 'H': 10, 'I': 10, 'J': 3, 'K': 36, 'L': 16, 'M': 30})
ws['A6'] = 'Datos'; ws['A6'].font = Font(bold=True)
inp(ws, f'A7:A{6 + N}')
D = f'$A$7:$A${6 + N}'
head(ws, 6, 4, ['xᵢ', 'nᵢ', 'fᵢ', 'Nᵢ', 'Fᵢ', '%'])
for i in range(40):
    r = 7 + i
    if i == 0:
        ws[f'D{r}'] = f'=IF(COUNT({D})=0,"",MIN({D}))'
    else:
        ws[f'D{r}'] = f'=IF(D{r-1}="","",IF(D{r-1}+1>MAX({D}),"",D{r-1}+1))'
    ws[f'D{r}'].fill = F_IN
    ws[f'E{r}'] = f'=IF(D{r}="","",COUNTIF({D},D{r}))'
    ws[f'F{r}'] = f'=IF(D{r}="","",E{r}/COUNT({D}))'
    ws[f'G{r}'] = f'=IF(D{r}="","",SUM($E$7:E{r}))'
    ws[f'H{r}'] = f'=IF(D{r}="","",G{r}/COUNT({D}))'
    ws[f'I{r}'] = f'=IF(D{r}="","",F{r}*100)'
    for c in 'EFGHI':
        ws[f'{c}{r}'].fill = F_OUT; ws[f'{c}{r}'].number_format = '0.0000' if c in 'FH' else ('0.00' if c == 'I' else '0')
ws['L5'] = 0.9; ws['L5'].fill = F_IN; ws['K5'] = 'Percentil que quieres (en tanto por uno) →'; ws['K5'].font = Font(bold=True, color='6B5800')
ws['L4'] = 2; ws['L4'].fill = F_IN; ws['K4'] = 'Valor k para "cuántos más de k" →'; ws['K4'].font = Font(bold=True, color='6B5800')
r = 7
outs = [('n', f'=COUNT({D})', '0'), ('Media', f'=AVERAGE({D})', None), ('Mediana', f'=MEDIAN({D})', None), ('Moda', f'=_xlfn.MODE.SNGL({D})', None),
        ('Mínimo', f'=MIN({D})', None), ('Máximo', f'=MAX({D})', None), ('Rango', f'=MAX({D})-MIN({D})', None),
        ('C₁ (CUARTIL.INC)', f'=_xlfn.QUARTILE.INC({D},1)', None), ('C₃ (CUARTIL.INC)', f'=_xlfn.QUARTILE.INC({D},3)', None),
        ('IQR = C₃ − C₁', '=L15-L14', None), ('Percentil elegido (celda L5)', f'=_xlfn.PERCENTILE.INC({D},L5)', None),
        ('Varianza S² (VAR.P)', f'=_xlfn.VAR.P({D})', None), ('Cuasivarianza s² (VAR.S)', f'=_xlfn.VAR.S({D})', None),
        ('Desviación típica S', f'=_xlfn.STDEV.P({D})', None), ('Cuasidesviación s', f'=_xlfn.STDEV.S({D})', None),
        ('Coeficiente de variación', f'=_xlfn.STDEV.P({D})/AVERAGE({D})', None), ('Asimetría (signo)', f'=_xlfn.SKEW.P({D})', None),
        ('Curtosis (signo)', f'=KURT({D})', None), ('Límite inferior atípicos', '=L14-1.5*L16', None), ('Límite superior atípicos', '=L15+1.5*L16', None),
        ('Nº de datos atípicos', f'=COUNTIF({D},"<"&L25)+COUNTIF({D},">"&L26)', '0'),
        ('Cuántos son MÁS de k (L4)', f'=COUNTIF({D},">"&L4)', '0'), ('Cuántos son k O MÁS', f'=COUNTIF({D},">="&L4)', '0'),
        ('Proporción con más de k', f'=COUNTIF({D},">"&L4)/COUNT({D})', None)]
for lab, f, fmt in outs:
    out(ws, r, 11, lab, f, fmt or '0.0000'); r += 1
ws['K32'] = 'Los cuartiles con la receta de clase (tabla) están en la plantilla B.'; ws['K32'].font = Font(italic=True, size=9, color='7A878E')

# ======================================================= B_Tabla_valores
ws = sheet('B_Tabla_valores', 'B · Tabla de valores y frecuencias', 't1', 'Cuando te dan la tabla (valores xᵢ y sus frecuencias nᵢ) en vez de los datos sueltos.',
           'Escribe los valores ORDENADOS de menor a mayor en A y sus frecuencias en B (absolutas o relativas). Usa la receta de clase para mediana y cuantiles.')
widths(ws, {'A': 10, 'B': 10, 'C': 10, 'D': 10, 'E': 12, 'F': 12, 'G': 3, 'H': 40, 'I': 16, 'J': 34})
head(ws, 6, 1, ['xᵢ', 'nᵢ', 'fᵢ', 'Nᵢ', 'xᵢ·nᵢ', 'xᵢ²·nᵢ'])
inp(ws, 'A7:B46')
for i in range(40):
    r = 7 + i
    ws[f'C{r}'] = f'=IF(B{r}="","",B{r}/SUM($B$7:$B$46))'
    ws[f'D{r}'] = f'=IF(B{r}="","",SUM($B$7:B{r}))'
    ws[f'E{r}'] = f'=IF(B{r}="","",A{r}*B{r})'
    ws[f'F{r}'] = f'=IF(B{r}="","",A{r}^2*B{r})'
    for c in 'CDEF':
        ws[f'{c}{r}'].fill = F_OUT; ws[f'{c}{r}'].number_format = '0.0000' if c == 'C' else 'General'
ws['I5'] = 0.75; ws['I5'].fill = F_IN; ws['H5'] = 'Cuantil que quieres (tanto por uno) →'; ws['H5'].font = Font(bold=True, color='6B5800')
n = 'SUM($B$7:$B$46)'; Nr = '$D$7:$D$46'; Xr = '$A$7:$A$46'


def receta(p):
    k = f'(COUNTIF({Nr},"<"&{p}*{n}-1E-9)+1)'
    return (f'=IF(ABS(INDEX({Nr},{k})-{p}*{n})<1E-9,(INDEX({Xr},{k})+INDEX({Xr},{k}+1))/2,INDEX({Xr},{k}))')


r = 7
for lab, f, fmt in [('n (suma de frecuencias)', f'={n}', 'General'), ('Media', f'=SUM(E7:E46)/{n}', None),
                    ('Varianza S²', f'=SUM(F7:F46)/{n}-(SUM(E7:E46)/{n})^2', None), ('Desviación típica S', '=SQRT(I9)', None),
                    ('Cuasivarianza s² (solo con frecuencias absolutas)', f'=I9*{n}/({n}-1)', None), ('Coeficiente de variación', '=I10/I8', None),
                    ('Moda', f'=INDEX({Xr},MATCH(MAX($B$7:$B$46),$B$7:$B$46,0))', 'General'),
                    ('Mediana (receta de clase)', receta('0.5'), 'General'), ('C₁ (receta de clase)', receta('0.25'), 'General'),
                    ('C₃ (receta de clase)', receta('0.75'), 'General'), ('Cuantil elegido (I5)', receta('I5'), 'General')]:
    out(ws, r, 8, lab, f, fmt or '0.0000'); r += 1

# ======================================================= C_Intervalos
ws = sheet('C_Intervalos', 'C · Datos agrupados en intervalos', 't1', 'Escribe el extremo inferior, el superior y la frecuencia de cada intervalo (en orden).',
           'Sirve con amplitudes iguales o distintas: la moda usa densidades. Si te dan %, escríbelos como frecuencias (la media sale igual).')
widths(ws, {'A': 9, 'B': 9, 'C': 9, 'D': 10, 'E': 10, 'F': 10, 'G': 12, 'H': 14, 'I': 10, 'J': 10, 'K': 3, 'L': 40, 'M': 16})
head(ws, 6, 1, ['L inf', 'L sup', 'nᵢ', 'Marca xᵢ', 'Amplitud', 'Densidad', 'xᵢ·nᵢ', 'xᵢ²·nᵢ', 'Nᵢ', 'fᵢ'])
inp(ws, 'A7:C21')
for i in range(15):
    r = 7 + i
    ws[f'D{r}'] = f'=IF(C{r}="","",(A{r}+B{r})/2)'
    ws[f'E{r}'] = f'=IF(C{r}="","",B{r}-A{r})'
    ws[f'F{r}'] = f'=IF(C{r}="","",C{r}/E{r})'
    ws[f'G{r}'] = f'=IF(C{r}="","",D{r}*C{r})'
    ws[f'H{r}'] = f'=IF(C{r}="","",D{r}^2*C{r})'
    ws[f'I{r}'] = f'=IF(C{r}="","",SUM($C$7:C{r}))'
    ws[f'J{r}'] = f'=IF(C{r}="","",C{r}/SUM($C$7:$C$21))'
    for c in 'DEFGHIJ':
        ws[f'{c}{r}'].fill = F_OUT; ws[f'{c}{r}'].number_format = 'General'
ws['M5'] = 0.9; ws['M5'].fill = F_IN; ws['L5'] = 'Percentil que quieres (tanto por uno) →'; ws['L5'].font = Font(bold=True, color='6B5800')
n = 'SUM($C$7:$C$21)'; Nr = '$I$7:$I$21'; Lr = '$A$7:$A$21'; Cr = '$C$7:$C$21'; Ar = '$E$7:$E$21'; Dr = '$F$7:$F$21'; last = 'COUNT($C$7:$C$21)'


def pq(p):
    k = f'(COUNTIF({Nr},"<"&{p}*{n}-1E-9)+1)'
    prev = f'IF({k}=1,0,INDEX({Nr},{k}-1))'
    return f'=INDEX({Lr},{k})+({p}*{n}-{prev})/INDEX({Cr},{k})*INDEX({Ar},{k})'


km = f'MATCH(MAX({Dr}),{Dr},0)'
moda = (f'=INDEX({Lr},{km})+(INDEX({Dr},{km})-IF({km}=1,0,INDEX({Dr},{km}-1)))/((INDEX({Dr},{km})-IF({km}=1,0,INDEX({Dr},{km}-1)))'
        f'+(INDEX({Dr},{km})-IF({km}={last},0,INDEX({Dr},{km}+1))))*INDEX({Ar},{km})')
r = 7
for lab, f in [('n', f'={n}'), ('Media', f'=SUM(G7:G21)/{n}'), ('Varianza S²', f'=SUM(H7:H21)/{n}-M8^2'), ('Desviación típica S', '=SQRT(M9)'),
               ('Cuasivarianza s²', f'=M9*{n}/({n}-1)'), ('Coeficiente de variación', '=M10/M8'),
               ('Intervalo modal (fila de la tabla)', f'={km}'), ('Moda', moda), ('Mediana', pq('0.5')), ('C₁', pq('0.25')), ('C₃', pq('0.75')),
               ('IQR', '=M17-M16'), ('Percentil elegido (M5)', pq('M5'))]:
    out(ws, r, 12, lab, f, 'General' if lab in ('n', 'Intervalo modal (fila de la tabla)') else '0.0000'); r += 1

# ======================================================= D_Grupos
ws = sheet('D_Grupos', 'D · Grupos: media global y varianza dentro / entre', 't1', 'Para cada grupo: tamaño, media y desviación típica O varianza (rellena una de las dos).',
           'Si te dan la desviación típica, déjala en la columna D y la varianza se calcula sola.')
widths(ws, {'A': 16, 'B': 10, 'C': 12, 'D': 12, 'E': 12, 'F': 12, 'G': 16, 'H': 10, 'I': 3, 'J': 40, 'K': 16})
head(ws, 6, 1, ['Grupo', 'Tamaño T', 'Media', 'Desv. típica', 'Varianza (si te la dan)', 'Varianza usada', '(media−global)²·T', 'CV'])
inp(ws, 'A7:E14')
for i in range(8):
    r = 7 + i
    ws[f'F{r}'] = f'=IF(B{r}="","",IF(E{r}<>"",E{r},D{r}^2))'
    ws[f'G{r}'] = f'=IF(B{r}="","",(C{r}-$K$7)^2*B{r})'
    ws[f'H{r}'] = f'=IF(B{r}="","",SQRT(F{r})/C{r})'
    for c in 'FGH':
        ws[f'{c}{r}'].fill = F_OUT; ws[f'{c}{r}'].number_format = '0.0000'
r = 7
for lab, f in [('Media global', '=SUMPRODUCT(B7:B14,C7:C14)/SUM(B7:B14)'), ('Varianza dentro', '=SUMPRODUCT(B7:B14,F7:F14)/SUM(B7:B14)'),
               ('Varianza entre', '=SUM(G7:G14)/SUM(B7:B14)'), ('Varianza total', '=K8+K9'), ('Desviación típica total', '=SQRT(K10)'),
               ('% debido a diferencias entre grupos', '=K9/K10'), ('Grupo más homogéneo (menor CV)', '=INDEX(A7:A14,MATCH(MIN(H7:H14),H7:H14,0))')]:
    out(ws, r, 10, lab, f, '0.00%' if '%' in lab else ('General' if 'homogéneo' in lab else '0.0000')); r += 1

# ======================================================= E_Transformacion
ws = sheet('E_Transformacion', 'E · Cambios de origen y escala (Y = a·X + b)', 't1', 'Qué pasa con la media y la varianza si todos los datos se multiplican por a y se les suma b.',
           'Subir un 3 % → a = 1,03. Bajar un 2 % → a = 0,98. Sumar 50 € → b = 50. Si solo sumas, a = 1. Si solo multiplicas, b = 0.')
widths(ws, {'A': 40, 'B': 16, 'C': 50})
for i, (lab, v) in enumerate([('Media original', 2550), ('Varianza original', 136500), ('a (multiplica)', 1.01), ('b (suma)', 25)]):
    ws.cell(6 + i, 1, lab).font = Font(bold=True); c = ws.cell(6 + i, 2, v); c.fill = F_IN; c.border = BOX
r = 11
for lab, f, note in [('Nueva media = a·media + b', '=B8*B6+B9', ''), ('Nueva varianza = a²·varianza', '=B8^2*B7', 'Sumar (b) no cambia la varianza'),
                     ('Nueva desviación típica = |a|·S', '=ABS(B8)*SQRT(B7)', ''), ('CV antes', '=SQRT(B7)/B6', ''), ('CV después', '=B13/B11', 'Solo multiplicar no cambia el CV')]:
    out(ws, r, 1, lab, f, '0.0000', note); r += 1

# ======================================================= F_Doble_entrada
ws = sheet('F_Doble_entrada', 'F · Tabla de doble entrada', 't2', 'Escribe los nombres de filas (A8:A13) y columnas (B7:G7) y las frecuencias.',
           'Calcula totales, frecuencias relativas, la tabla de independencia, las condicionadas y si son independientes.')
widths(ws, {c: 12 for c in 'ABCDEFGHIJKLMNOP'}); ws.column_dimensions['A'].width = 16
ws['A7'] = 'X \\ Y'; ws['A7'].font = Font(bold=True)
inp(ws, 'B7:G7'); inp(ws, 'A8:A13'); inp(ws, 'B8:G13')
ws['H7'] = 'Total'; ws['H7'].font = Font(bold=True); ws['A14'] = 'Total'; ws['A14'].font = Font(bold=True)
for i in range(6):
    r = 8 + i
    ws[f'H{r}'] = f'=IF(COUNT(B{r}:G{r})=0,"",SUM(B{r}:G{r}))'; ws[f'H{r}'].fill = F_OUT
for j in range(7):
    c = CL(2 + j)
    ws[f'{c}14'] = f'=IF(COUNT({c}8:{c}13)=0,"",SUM({c}8:{c}13))'; ws[f'{c}14'].fill = F_OUT


def block(r0, title, fn, fmt='0.0000'):
    ws.cell(r0, 1, title).font = Font(bold=True, color=TEMA['t2'])
    for j in range(6):
        c = CL(2 + j); ws[f'{c}{r0 + 1}'] = f'=IF({c}$7="","",{c}$7)'; ws[f'{c}{r0 + 1}'].font = Font(bold=True)
    for i in range(6):
        r = r0 + 2 + i; ws[f'A{r}'] = f'=IF($A{8 + i}="","",$A{8 + i})'; ws[f'A{r}'].font = Font(bold=True)
        for j in range(6):
            c = CL(2 + j); ws[f'{c}{r}'] = fn(c, 8 + i); ws[f'{c}{r}'].fill = F_OUT; ws[f'{c}{r}'].number_format = fmt


block(16, 'Frecuencias relativas fᵢⱼ = nᵢⱼ / n', lambda c, r: f'=IF({c}{r}="","",{c}{r}/$H$14)')
block(25, 'Si fueran independientes: fᵢ. × f.ⱼ', lambda c, r: f'=IF({c}{r}="","",($H{r}/$H$14)*({c}$14/$H$14))')
block(34, 'Condicionadas por filas (Y | X): % de cada fila', lambda c, r: f'=IF({c}{r}="","",{c}{r}/$H{r})', '0.0%')
block(43, 'Condicionadas por columnas (X | Y): % de cada columna', lambda c, r: f'=IF({c}{r}="","",{c}{r}/{c}$14)', '0.0%')
block(52, 'Diferencias |fᵢⱼ − fᵢ.·f.ⱼ|', lambda c, r: f'=IF({c}{r + 10}="","",ABS({c}{r + 10}-{c}{r + 19}))')
out(ws, 61, 1, 'Mayor diferencia', '=MAX(B54:G59)', '0.0000')
out(ws, 62, 1, '¿Independientes?', '=IF(B61<0.0005,"Sí: coinciden en todas las casillas","No: hay casillas donde fᵢⱼ ≠ fᵢ.·f.ⱼ")', 'General')
ws.column_dimensions['B'].width = 14

# ======================================================= G_Regresion
ws = sheet('G_Regresion', 'G · Regresión y correlación con datos', 't3', 'Pega X en la columna A y Y en la B (desde la fila 7). Hasta 500 pares.',
           'La recta de Y sobre X es la que se usa para predecir Y. Para la nube de puntos: selecciona A y B → Insertar → Dispersión → línea de tendencia.')
widths(ws, {'A': 10, 'B': 10, 'C': 3, 'D': 40, 'E': 16, 'F': 36})
head(ws, 6, 1, ['X', 'Y'])
inp(ws, f'A7:B{6 + N}')
X = f'$A$7:$A${6 + N}'; Y = f'$B$7:$B${6 + N}'
ws['E5'] = 11; ws['E5'].fill = F_IN; ws['D5'] = 'Valor de X para la predicción →'; ws['D5'].font = Font(bold=True, color='6B5800')
r = 7
for lab, f in [('n', f'=COUNT({X})'), ('Media de X', f'=AVERAGE({X})'), ('Media de Y', f'=AVERAGE({Y})'), ('Varianza de X', f'=_xlfn.VAR.P({X})'),
               ('Varianza de Y', f'=_xlfn.VAR.P({Y})'), ('Covarianza', f'=_xlfn.COVARIANCE.P({X},{Y})'), ('Pendiente b (Y sobre X)', f'=SLOPE({Y},{X})'),
               ('Término independiente a', f'=INTERCEPT({Y},{X})'), ('Correlación r', f'=CORREL({X},{Y})'), ('R²', f'=RSQ({Y},{X})'),
               ('1 − R² (no explicado)', '=1-E16'), ('Predicción de Y para X (E5)', f'=_xlfn.FORECAST.LINEAR(E5,{Y},{X})'),
               ("Pendiente b' (X sobre Y)", f'=SLOPE({X},{Y})'), ("Término a' (X sobre Y)", f'=INTERCEPT({X},{Y})')]:
    out(ws, r, 4, lab, f, 'General' if lab == 'n' else '0.0000'); r += 1
ws['F13'] = 'Por cada unidad más de X, Y cambia b unidades de media'; ws['F13'].font = Font(italic=True, size=9, color='7A878E')

# ======================================================= H_Reg_resumen
ws = sheet('H_Reg_resumen', 'H · Regresión a partir de resúmenes', 't3', 'Cuando no te dan los datos sino medias, desviaciones y covarianza (o r).',
           'Rellena una opción de cada par (por ejemplo, Sx O Sx²). Si te dan r en vez de Sxy, deja Sxy vacío.')
widths(ws, {'A': 34, 'B': 16, 'C': 4, 'D': 40, 'E': 16})
for i, (lab, v) in enumerate([('x̄', 60), ('ȳ', 850), ('Sx (desv. típica de X)', 40), ('Sx² (varianza de X)', None), ('Sy (desv. típica de Y)', 500),
                               ('Sy² (varianza de Y)', None), ('Sxy (covarianza)', 19200), ('r (correlación)', None), ('X para predecir', 100)]):
    ws.cell(6 + i, 1, lab).font = Font(bold=True); c = ws.cell(6 + i, 2, v); c.fill = F_IN; c.border = BOX
r = 6
for lab, f in [('Sx² usada', '=IF(B9<>"",B9,B8^2)'), ('Sy² usada', '=IF(B11<>"",B11,B10^2)'), ('Sxy usada', '=IF(B12<>"",B12,B13*SQRT(E6)*SQRT(E7))'),
               ('Pendiente b = Sxy/Sx²', '=E8/E6'), ('a = ȳ − b·x̄', '=B7-E9*B6'), ('r = Sxy/(Sx·Sy)', '=E8/SQRT(E6*E7)'), ('R²', '=E11^2'),
               ('Predicción Ŷ', '=E10+E9*B14'), ("b' (X sobre Y) = Sxy/Sy²", '=E8/E7'), ("a' = x̄ − b'·ȳ", '=B6-E14*B7')]:
    out(ws, r, 4, lab, f); r += 1

# ======================================================= I_Indices
ws = sheet('I_Indices', 'I · Serie temporal: índices y tasas', 't4', 'Años en A y valores en B (desde la fila 7, en orden). Elige el año base en F5.',
           'Calcula índices con la base elegida, tasas año a año, tasa total y tasa media anual.')
widths(ws, {'A': 10, 'B': 12, 'C': 14, 'D': 14, 'E': 3, 'F': 38, 'G': 16})
head(ws, 6, 1, ['Año', 'Valor', 'Índice (base)', 'Tasa vs anterior'])
inp(ws, 'A7:B36')
ws['G5'] = 2015; ws['G5'].fill = F_IN; ws['F5'] = 'Año base →'; ws['F5'].font = Font(bold=True, color='6B5800')
for i in range(30):
    r = 7 + i
    ws[f'C{r}'] = f'=IF(B{r}="","",B{r}/INDEX($B$7:$B$36,MATCH($G$5,$A$7:$A$36,0))*100)'
    ws[f'D{r}'] = '' if i == 0 else f'=IF(B{r}="","",B{r}/B{r-1}-1)'
    ws[f'C{r}'].fill = F_OUT; ws[f'D{r}'].fill = F_OUT; ws[f'C{r}'].number_format = '0.00'; ws[f'D{r}'].number_format = '0.00%'
r = 7
for lab, f, fmt in [('Primer año', '=A7', '0'), ('Último año', '=INDEX(A7:A36,COUNT(B7:B36))', '0'), ('Años transcurridos', '=G8-G7', '0'),
                    ('Tasa total (primero → último)', '=INDEX(B7:B36,COUNT(B7:B36))/B7-1', '0.00%'), ('Tasa media anual', '=(1+G10)^(1/G9)-1', '0.000%')]:
    out(ws, r, 6, lab, f, fmt); r += 1

# ======================================================= J_Tasas
ws = sheet('J_Tasas', 'J · Tasas encadenadas y equivalentes', 't4', 'Escribe las tasas de cada periodo en B (como 5 % o 0,05; negativas para bajadas).',
           'Calcula los índices encadenados, la tasa acumulada y la tasa media. Abajo: de anual a mensual y al revés.')
widths(ws, {'A': 14, 'B': 12, 'C': 16, 'D': 3, 'E': 38, 'F': 16})
head(ws, 6, 1, ['Periodo', 'Tasa', 'Índice (inicio = 100)'])
ws['A7'] = 'Inicio'; ws['C7'] = 100; ws['C7'].fill = F_OUT
inp(ws, 'A8:B31')
for i in range(24):
    r = 8 + i
    ws[f'B{r}'].number_format = '0.00%'
    ws[f'C{r}'] = f'=IF(B{r}="","",C{r-1}*(1+B{r}))'; ws[f'C{r}'].fill = F_OUT; ws[f'C{r}'].number_format = '0.00'
r = 7
for lab, f, fmt in [('Nº de periodos', '=COUNT(B8:B31)', '0'), ('Tasa acumulada', '=INDEX(C7:C31,F7+1)/100-1', '0.00%'),
                    ('Tasa media por periodo (geométrica)', '=(1+F8)^(1/F7)-1', '0.000%')]:
    out(ws, r, 5, lab, f, fmt); r += 1
ws['E12'] = 'Tasa anual →'; ws['E12'].font = Font(bold=True, color='6B5800'); ws['F12'] = 0.06; ws['F12'].fill = F_IN; ws['F12'].number_format = '0.00%'
out(ws, 13, 5, 'Tasa mensual equivalente', '=(1+F12)^(1/12)-1', '0.0000%')
ws['E15'] = 'Tasa mensual →'; ws['E15'].font = Font(bold=True, color='6B5800'); ws['F15'] = 0.005; ws['F15'].fill = F_IN; ws['F15'].number_format = '0.00%'
out(ws, 16, 5, 'Tasa anual equivalente', '=(1+F15)^12-1', '0.00%')

# ======================================================= K_Laspeyres
ws = sheet('K_Laspeyres', 'K · Laspeyres, Paasche y Fisher', 't4', 'Por producto: precio y cantidad en el año base (0) y en el año actual (t).',
           'Índices de precios y de cantidades.')
widths(ws, {'A': 14, 'B': 10, 'C': 10, 'D': 10, 'E': 10, 'F': 11, 'G': 11, 'H': 11, 'I': 11, 'J': 11, 'K': 11, 'L': 3, 'M': 36, 'N': 14})
head(ws, 6, 1, ['Producto', 'p₀', 'q₀', 'pₜ', 'qₜ', 'pₜ·q₀', 'p₀·q₀', 'pₜ·qₜ', 'p₀·qₜ', 'qₜ·p₀ ', 'q₀·pₜ'])
inp(ws, 'A7:E26')
for i in range(20):
    r = 7 + i
    for c, f in zip('FGHIJK', [f'D{r}*C{r}', f'B{r}*C{r}', f'D{r}*E{r}', f'B{r}*E{r}', f'E{r}*B{r}', f'C{r}*D{r}']):
        ws[f'{c}{r}'] = f'=IF(B{r}="","",{f})'; ws[f'{c}{r}'].fill = F_OUT
r = 7
for lab, f in [('Laspeyres de precios', '=SUM(F7:F26)/SUM(G7:G26)*100'), ('Paasche de precios', '=SUM(H7:H26)/SUM(I7:I26)*100'),
               ('Fisher de precios', '=SQRT(N7*N8)'), ('Laspeyres de cantidades', '=SUM(J7:J26)/SUM(G7:G26)*100'),
               ('Paasche de cantidades', '=SUM(H7:H26)/SUM(K7:K26)*100'), ('Fisher de cantidades', '=SQRT(N10*N11)')]:
    out(ws, r, 13, lab, f, '0.00'); r += 1

# ======================================================= L_Deflactar
ws = sheet('L_Deflactar', 'L · Deflactar con el IPC (valores reales)', 't4', 'Años en A, valores nominales (salario, alquiler…) en B e IPC en C. Elige el año de referencia en H5.',
           'Valor real = nominal / IPC × IPC del año de referencia. Sale en euros del año de referencia.')
widths(ws, {'A': 10, 'B': 12, 'C': 10, 'D': 14, 'E': 14, 'F': 14, 'G': 3, 'H': 40, 'I': 16})
head(ws, 6, 1, ['Año', 'Nominal', 'IPC', 'Real (€ año ref.)', 'Índice nominal', 'Índice real'])
inp(ws, 'A7:C36')
ws['I5'] = 2015; ws['I5'].fill = F_IN; ws['H5'] = 'Año de referencia →'; ws['H5'].font = Font(bold=True, color='6B5800')
ipc0 = 'INDEX($C$7:$C$36,MATCH($I$5,$A$7:$A$36,0))'; nom0 = 'INDEX($B$7:$B$36,MATCH($I$5,$A$7:$A$36,0))'
for i in range(30):
    r = 7 + i
    ws[f'D{r}'] = f'=IF(B{r}="","",B{r}/C{r}*{ipc0})'
    ws[f'E{r}'] = f'=IF(B{r}="","",B{r}/{nom0}*100)'
    ws[f'F{r}'] = f'=IF(B{r}="","",D{r}/INDEX($D$7:$D$36,MATCH($I$5,$A$7:$A$36,0))*100)'
    for c in 'DEF':
        ws[f'{c}{r}'].fill = F_OUT; ws[f'{c}{r}'].number_format = '0.00'
r = 7
for lab, f, fmt in [('Variación nominal (primero → último)', '=INDEX(B7:B36,COUNT(B7:B36))/B7-1', '0.00%'),
                    ('Variación real (primero → último)', '=INDEX(D7:D36,COUNT(B7:B36))/D7-1', '0.00%'),
                    ('Inflación del periodo', '=INDEX(C7:C36,COUNT(B7:B36))/C7-1', '0.00%'),
                    ('Para mantener el poder adquisitivo, el último año debería valer', '=B7*INDEX(C7:C36,COUNT(B7:B36))/C7', '0.00')]:
    out(ws, r, 8, lab, f, fmt); r += 1

# ======================================================= Portada
ws = wb.create_sheet('Inicio', 0)
ws.sheet_view.showGridLines = False
widths(ws, {'A': 3, 'B': 22, 'C': 54, 'D': 60})
ws['B2'] = 'Plantillas de estadística'; ws['B2'].font = Font(bold=True, size=18, color=INK)
lines = ['Calculadoras listas: pegas los datos del ejercicio en las celdas amarillas y los resultados salen en las verdes.',
         'Úsalas para COMPROBAR lo que haces tú. En el examen, solo si la profesora permite llevar archivos propios.',
         'Para cada hoja: primero lee la frase de la fila 3. Si algo sale #N/A o #¡DIV/0!, revisa que los datos estén en las columnas correctas.',
         'Empieza por la hoja «Traductor»: te dice qué plantilla usar según lo que diga el enunciado.']
for i, t in enumerate(lines):
    ws.cell(4 + i, 2, t).font = Font(size=11)
head(ws, 9, 2, ['Hoja', 'Para qué', 'Cómo'])
items = [('Traductor', '¿Qué me están pidiendo?', 'Frases típicas del enunciado → qué calcular')] + TOC
for i, (nm, t, d) in enumerate(items):
    r = 10 + i
    c = ws.cell(r, 2, nm); c.hyperlink = f"#'{nm}'!A1"; c.font = Font(color='0563C1', underline='single')
    ws.cell(r, 3, t); ws.cell(r, 4, d).alignment = Alignment(wrap_text=True)
for w in wb.worksheets:
    w.freeze_panes = w.freeze_panes or None
wb.save('/home/claude/xl2/Plantillas_estadistica.xlsx')
print('ok', wb.sheetnames)
