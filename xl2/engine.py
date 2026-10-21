"""Motor de autocorrección para libros de Excel (openpyxl)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

INK = '1F2A30'
F_DATA = PatternFill('solid', fgColor='E6EDFC')    # datos (azul claro)
F_IN = PatternFill('solid', fgColor='FFF4C2')      # tu respuesta (amarillo)
F_HEAD = PatternFill('solid', fgColor='DDE5EA')
F_OK = PatternFill('solid', fgColor='C8EFD9')
F_BAD = PatternFill('solid', fgColor='FAD2D8')
F_TITLE = PatternFill('solid', fgColor='1F2A30')
thin = Side(style='thin', color='B7C2C8')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TEMA = {'t1': '2F5FD0', 't2': '0F8A68', 't3': 'C2185B', 't4': 'A86F00', 'tx': '5B6B78'}


def tol_for(key, dec):
    return max(0.5 * 10 ** (-dec) * 1.02, abs(key) * 0.001)


def xlnum(v):
    """número para incrustar en una fórmula (formato inglés de openpyxl)"""
    return repr(float(v)) if not float(v).is_integer() else str(int(v))


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.claves = self.wb.create_sheet('Claves')
        self.claves.sheet_state = 'veryHidden'
        self.kn = 0
        self.registry = []   # para pruebas: (hoja, celda, clave, alternativas, tipo)
        self.sheets = []

    def key(self, value):
        self.kn += 1
        self.claves.cell(self.kn, 1, value)
        return f"Claves!$A${self.kn}"

    def sheet(self, name, title, color='tx', subtitle=None):
        s = Sheet(self, name, title, color, subtitle)
        self.sheets.append(s)
        return s

    def save(self, path):
        # la hoja de claves al final
        self.wb.move_sheet('Claves', offset=len(self.wb.sheetnames))
        self.wb.save(path)


class Sheet:
    def __init__(self, book, name, title, color, subtitle):
        self.b = book
        self.ws = book.wb.create_sheet(name[:31])
        self.name = name[:31]
        self.color = TEMA.get(color, color)
        ws = self.ws
        ws.sheet_view.showGridLines = True
        ws.sheet_view.zoomScale = 115
        ws.sheet_properties.tabColor = self.color
        widths = {'A': 42, 'B': 14, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 13, 'K': 13, 'L': 13}
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        c = ws.cell(1, 1, title)
        c.font = Font(bold=True, size=15, color='FFFFFF')
        for col in range(1, 8):
            ws.cell(1, col).fill = PatternFill('solid', fgColor=self.color)
        ws.row_dimensions[1].height = 30
        c.alignment = Alignment(vertical='center')
        self.r = 2
        if subtitle:
            c = ws.cell(2, 1, subtitle)
            c.font = Font(italic=True, color='4F5D65')
            self.r = 3
        # marcador de aciertos
        self.score_row = 2
        self.n_answers = 0
        self.r += 1
        self.legend()

    def legend(self):
        ws = self.ws
        r = self.r
        ws.cell(r, 1, 'Celdas azules: datos.  Celdas amarillas: tu respuesta (se ponen verdes si está bien y rojas si no).').font = Font(size=9, color='4F5D65')
        self.r += 2

    # ---------- texto ----------
    def h(self, text):
        c = self.ws.cell(self.r, 1, text)
        c.font = Font(bold=True, size=12, color=self.color)
        self.r += 1

    def p(self, text, italic=False, color='26323A', bold=False, size=11):
        c = self.ws.cell(self.r, 1, text)
        c.font = Font(italic=italic, color=color, bold=bold, size=size)
        c.alignment = Alignment(wrap_text=False)
        self.r += 1

    def hint(self, text):
        c = self.ws.cell(self.r, 1, '💡 ' + text)
        c.font = Font(italic=True, color='6B5800', size=10)
        self.r += 1

    def gap(self, n=1):
        self.r += n

    # ---------- datos ----------
    def data_row(self, values, per_row=10, label='Datos'):
        """escribe datos sueltos en filas de 10; devuelve el rango"""
        ws = self.ws
        ws.cell(self.r, 1, label).font = Font(bold=True)
        r0 = self.r
        for i, v in enumerate(values):
            rr = r0 + i // per_row
            cc = 2 + i % per_row
            c = ws.cell(rr, cc, v)
            c.fill = F_DATA
            c.border = BOX
            c.alignment = Alignment(horizontal='center')
        r1 = r0 + (len(values) - 1) // per_row
        self.r = r1 + 2
        last_col = 2 + min(per_row, len(values)) - 1
        return f"$B${r0}:${CL(last_col)}${r1}"

    def data_table(self, headers, rows, col=1):
        """tabla de datos (azul). Devuelve (fila_primera, fila_última)"""
        ws = self.ws
        for j, hd in enumerate(headers):
            c = ws.cell(self.r, col + j, hd)
            c.font = Font(bold=True, color=INK)
            c.fill = F_HEAD
            c.border = BOX
            c.alignment = Alignment(horizontal='center', wrap_text=True)
        r1 = self.r + 1
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = ws.cell(r1 + i, col + j, v)
                c.border = BOX
                if v is not None and j > 0 or (j == 0 and isinstance(v, (int, float))):
                    c.fill = F_DATA
                c.alignment = Alignment(horizontal='center')
        self.r = r1 + len(rows) + 1
        return r1, r1 + len(rows) - 1

    def work_table(self, headers, given, keys, hints=None, dec=None, col=1):
        """Tabla de trabajo.
        headers: cabeceras; given: lista de filas con las columnas dadas (datos);
        keys: lista de filas con los valores correctos de las columnas a rellenar;
        hints: textos de pista por columna a rellenar (solo nivel 1)."""
        ws = self.ws
        ng = len(given[0])
        for j, hd in enumerate(headers):
            c = ws.cell(self.r, col + j, hd)
            c.font = Font(bold=True, color=INK)
            c.fill = F_HEAD
            c.border = BOX
            c.alignment = Alignment(horizontal='center', wrap_text=True)
        self.r += 1
        if hints:
            for j, t in enumerate(hints):
                c = ws.cell(self.r, col + ng + j, 'ej. ' + t)
                c.font = Font(italic=True, size=9, color='6B5800')
                c.alignment = Alignment(horizontal='center', wrap_text=True)
            ws.cell(self.r, col, 'Pista →').font = Font(italic=True, size=9, color='6B5800')
            ws.row_dimensions[self.r].height = 30
            self.r += 1
        r1 = self.r
        for i, (g, k) in enumerate(zip(given, keys)):
            for j, v in enumerate(g):
                c = ws.cell(r1 + i, col + j, v)
                c.fill = F_DATA
                c.border = BOX
                c.alignment = Alignment(horizontal='center')
            for j, kv in enumerate(k):
                cc = col + ng + j
                ref = f"{CL(cc)}{r1 + i}"
                c = ws.cell(r1 + i, cc)
                c.fill = F_IN
                c.border = BOX
                d = dec[j] if dec else 2
                self.cf_cell(ref, kv, d)
        self.r = r1 + len(given) + 1
        return r1

    def cf_cell(self, ref, keyval, dec=2):
        kref = self.b.key(keyval)
        t = xlnum(tol_for(keyval, dec))
        ok = f'AND({ref}<>"",ISNUMBER({ref}),ABS({ref}-{kref})<={t})'
        bad = f'AND({ref}<>"",OR(NOT(ISNUMBER({ref})),ABS({ref}-{kref})>{t}))'
        self.ws.conditional_formatting.add(ref, FormulaRule(formula=[ok], fill=F_OK))
        self.ws.conditional_formatting.add(ref, FormulaRule(formula=[bad], fill=F_BAD))
        self.b.registry.append((self.name, ref, keyval, [], 'cell', dec))

    # ---------- preguntas ----------
    def ask(self, label, key, dec=2, alts=(), kind='num', accept=None, note=None, unit=''):
        """Pregunta con respuesta autocorregida.
        alts: lista de (valor_erróneo_típico, pista)."""
        ws = self.ws
        r = self.r
        a = ws.cell(r, 1, label)
        a.alignment = Alignment(wrap_text=True, vertical='center')
        a.font = Font(bold=True)
        if len(label) > 48:
            ws.row_dimensions[r].height = 15 * (1 + len(label) // 46)
        inp = ws.cell(r, 2)
        inp.fill = F_IN
        inp.border = BOX
        inp.alignment = Alignment(horizontal='center', vertical='center')
        ref = f"B{r}"
        if kind != 'text':
            alts = [(v, h) for v, h in alts if abs(v - key) > tol_for(key, dec) * 1.5]
        alts = [(v, h.replace('"', '""')) for v, h in alts]
        if kind == 'text':
            acc = accept or [key]
            kref = self.b.key(key)
            ok = 'OR(' + ','.join(f'UPPER(TRIM({ref}))="{x.upper()}"' for x in acc) + ')'
            chain = '"✘ Revisa la respuesta."'
            for v, h in reversed(list(alts)):
                chain = f'IF(UPPER(TRIM({ref}))="{str(v).upper()}","✘ {h}",{chain})'
        else:
            kref = self.b.key(key)
            t = xlnum(tol_for(key, dec))
            if kind == 'pct':
                ok = f'AND(ISNUMBER({ref}),OR(ABS({ref}-{kref})<={t},ABS({ref}/100-{kref})<={t}))'
            else:
                ok = f'AND(ISNUMBER({ref}),ABS({ref}-{kref})<={t})'
            chain = '"✘ Revisa el cálculo."'
            for v, h in reversed(list(alts)):
                tv = xlnum(tol_for(v, dec))
                vv = xlnum(v)
                cond = f'ABS({ref}-{vv})<={tv}'
                if kind == 'pct':
                    cond = f'OR(ABS({ref}-{vv})<={tv},ABS({ref}/100-{vv})<={tv})'
                chain = f'IF({cond},"✘ {h}",{chain})'
            chain = f'IF(NOT(ISNUMBER({ref})),"✘ Escribe solo el número (sin texto ni %).",{chain})'
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
        msg = ws.cell(r, 3, f'=IF({ref}="","",IF({ok},"✔ Bien",{chain}))')
        msg.alignment = Alignment(wrap_text=True, vertical='center')
        msg.font = Font(bold=True)
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT($C{r},1)="✔"'], fill=F_OK))
        ws.conditional_formatting.add(ref, FormulaRule(formula=[f'LEFT($C{r},1)="✘"'], fill=F_BAD))
        ws.conditional_formatting.add(f'C{r}:I{r}', FormulaRule(formula=[f'LEFT($C{r},1)="✔"'], font=Font(bold=True, color='0F7A4E')))
        ws.conditional_formatting.add(f'C{r}:I{r}', FormulaRule(formula=[f'LEFT($C{r},1)="✘"'], font=Font(bold=True, color='B3123A')))
        fmt = {0: '0', 1: '0.0', 2: '0.00', 3: '0.000', 4: '0.0000'}.get(dec, '0.0000')
        if kind != 'text':
            inp.number_format = 'General'
        if note:
            c = ws.cell(r, 10, note)
            c.font = Font(italic=True, size=9, color='7A878E')
        self.b.registry.append((self.name, ref, key, [(v, h.replace('""', '"')) for v, h in alts], kind, dec))
        self.n_answers += 1
        self.r += 1
        return ref

    def finish(self):
        ws = self.ws
        r = self.score_row
        ws.cell(r, 5, 'Aciertos:').font = Font(bold=True)
        c = ws.cell(r, 6, f'=SUMPRODUCT(--(LEFT($C$1:$C${self.r + 5},1)="✔"))')
        c.font = Font(bold=True, size=13, color='0F7A4E')
        c.alignment = Alignment(horizontal='right')
        ws.cell(r, 7, f'de {self.n_answers}').font = Font(bold=True)
        ws.freeze_panes = 'A2'
        return f"'{self.name}'!$F${r}"
