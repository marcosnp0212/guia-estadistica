import sys
from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink
for p in sys.argv[1:]:
    wb=load_workbook(p);n=0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                h=c.hyperlink
                if h is not None and h.target and h.target.startswith('#'):
                    loc=h.target[1:];c.hyperlink=Hyperlink(ref=c.coordinate,location=loc);n+=1
    wb.save(p);print(p,n,[ws.title for ws in wb.worksheets if ws.sheet_state!='visible'])
