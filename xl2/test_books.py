import pickle,subprocess,os,shutil
from openpyxl import load_workbook
reg=pickle.load(open('registry.pkl','rb'))
def run(src,regs,scen):
    wb=load_workbook(src)
    exp={}
    for (sh,ref,key,alts,kind,dec) in regs:
        if kind=='cell': continue
        ws=wb[sh]
        if scen=='ok': v=round(key,dec) if kind!='text' else key; e='✔'
        elif scen=='pct':
            if kind!='pct': continue
            v=round(key*100,max(dec-2,0)+1); e='✔'
        elif scen=='alt':
            if not alts: continue
            v=alts[0][0]; v=round(v,dec) if kind!='text' else v; e='✘ '+alts[0][1]
        elif scen=='bad':
            v=(key*1.7+3.3) if kind!='text' else 'zzz'; e='✘'
        ws[ref]=v; exp[(sh,ref)]=e
    out=f'/tmp/t_{scen}_{os.path.basename(src)}';wb.save(out)
    shutil.rmtree('/tmp/rc',ignore_errors=True);os.makedirs('/tmp/rc')
    subprocess.run(['soffice','--headless','--calc','--convert-to','xlsx','--outdir','/tmp/rc',out],capture_output=True)
    wr=load_workbook('/tmp/rc/'+os.path.basename(out),data_only=True)
    bad=[]
    for (sh,ref),e in exp.items():
        r=int(ref[1:]);msg=wr[sh][f'C{r}'].value
        if msg is None or not str(msg).startswith(e[:60]): bad.append((sh,ref,wr[sh][ref].value,msg,e[:70]))
    return len(exp),bad,wr
for name,key in [('Practicas_autocorregidas.xlsx','p'),('Simulacros_autocorregidos.xlsx','s')]:
    for scen in ['ok','pct','alt','bad']:
        n,bad,wr=run(name,reg[key],scen)
        print(name,scen,n,'fallos:',len(bad))
        for b in bad[:8]: print('   ',b)
        if scen=='ok':
            ini=wr['Inicio'];vals=[ini.cell(r,c).value for r in range(1,80) for c in (3,4,5) if ini.cell(r,c).value is not None and ('de' in str(ini.cell(r,c).value) or isinstance(ini.cell(r,c).value,float))]
            print('   inicio:',vals[:12],'...',vals[-4:])
