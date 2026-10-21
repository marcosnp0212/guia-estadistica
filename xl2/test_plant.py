from openpyxl import load_workbook
import subprocess,shutil,os
wb=load_workbook('Plantillas_estadistica.xlsx')
herm=[0,2,0,0,2,2,0,0,1,1,3,0,0,1,0,0,1,0,1,4,0,0,1,4,2,0,0,1,0,0,2,2,1,1,0,0,0,5,1,3,0,1,0,1]
ws=wb['A_Discreta']
for i,v in enumerate(herm): ws.cell(7+i,1,v)
ws=wb['B_Tabla_valores']
for i,(x,n) in enumerate(zip(range(6),[21,12,6,2,2,1])): ws.cell(7+i,1,x);ws.cell(7+i,2,n)
ws=wb['C_Intervalos']
for i,(a,b,n) in enumerate([(160,170,3),(170,180,6),(180,190,5),(190,200,2)]): ws.cell(7+i,1,a);ws.cell(7+i,2,b);ws.cell(7+i,3,n)
ws['M5']=0.7
ws=wb['D_Grupos']
for i,(g,t,m,v) in enumerate([('A',150,10,10),('B',200,10,25),('C',150,12,16)]): ws.cell(7+i,1,g);ws.cell(7+i,2,t);ws.cell(7+i,3,m);ws.cell(7+i,5,v)
ws=wb['F_Doble_entrada']
for j,c in enumerate(['Presencial','Híbrido','Remoto']): ws.cell(7,2+j,c)
for i,(r,vals) in enumerate([('RRHH',[8,3,2]),('Finanzas',[5,4,1]),('Producción',[10,6,1])]):
    ws.cell(8+i,1,r)
    for j,v in enumerate(vals): ws.cell(8+i,2+j,v)
ws=wb['G_Regresion']
for i,(x,y) in enumerate(zip([1,2,3,5,6,7,9,10,12,15],[1400,1600,1700,1700,1800,1850,2200,2400,2250,2500])): ws.cell(7+i,1,x);ws.cell(7+i,2,y)
ws=wb['I_Indices']
for i,v in enumerate([1500,1530,1570,1600,1650,1700,1750,1800,1850,1900,1950]): ws.cell(7+i,1,2015+i);ws.cell(7+i,2,v)
ws['G5']=2019
ws=wb['J_Tasas']
for i,t in enumerate([.09,.10,-.02,.08,.10]): ws.cell(8+i,1,2021+i);ws.cell(8+i,2,t)
ws=wb['K_Laspeyres']
for i,row in enumerate([('Pan',1,100,1.2,110),('Leche',.9,200,1.1,190)]):
    for j,v in enumerate(row): ws.cell(7+i,1+j,v)
ws=wb['L_Deflactar']
for i,(v,c) in enumerate(zip([1500,1530,1570,1600,1650,1700,1750,1800,1850,1900,1950],[100,101.2,102.5,104,106,108.5,110.2,112,114.5,117,119])): ws.cell(7+i,1,2015+i);ws.cell(7+i,2,v);ws.cell(7+i,3,c)
wb.save('/tmp/plt.xlsx')
shutil.rmtree('/tmp/rc2',ignore_errors=True);os.makedirs('/tmp/rc2')
subprocess.run(['soffice','--headless','--calc','--convert-to','xlsx','--outdir','/tmp/rc2','/tmp/plt.xlsx'],capture_output=True)
w=load_workbook('/tmp/rc2/plt.xlsx',data_only=True)
def show(sh,col,r0,r1):
    ws=w[sh];print(sh,'|',' ; '.join(f'{ws.cell(r,col-1).value}={ws.cell(r,col).value if not isinstance(ws.cell(r,col).value,float) else round(ws.cell(r,col).value,4)}' for r in range(r0,r1+1)))
show('A_Discreta',12,7,30);print('  tabla',[w['A_Discreta'].cell(7+i,5).value for i in range(8)])
show('B_Tabla_valores',9,7,17);show('C_Intervalos',13,7,19);show('D_Grupos',11,7,13);show('E_Transformacion',2,11,15)
show('F_Doble_entrada',2,61,62);print('  indep row',[round(w['F_Doble_entrada'].cell(27,c).value,4) for c in range(2,5)])
show('G_Regresion',5,7,20);show('H_Reg_resumen',5,6,15);show('I_Indices',7,7,11);print('  idx',[round(w['I_Indices'].cell(7+i,3).value,2) for i in range(11)])
show('J_Tasas',6,7,9);show('J_Tasas',6,13,13);show('J_Tasas',6,16,16);show('K_Laspeyres',14,7,12);show('L_Deflactar',9,7,10);print('  real idx',round(w['L_Deflactar'].cell(17,6).value,2))
errs=[(s.title,c.coordinate,c.value) for s in w.worksheets for row in s.iter_rows() for c in row if isinstance(c.value,str) and c.value.startswith('#')]
print('errores',errs[:10],len(errs))
