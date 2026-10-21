import sys; sys.path.insert(0, '/home/claude/xl2')
import numpy as np
from engine import *
from tipos import *
from openpyxl.styles import Font, Alignment, PatternFill


def cover(book, title, lines, rows_title='Hoja'):
    ws = book.wb.create_sheet('Inicio', 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 62
    ws.column_dimensions['D'].width = 14
    ws['B2'] = title
    ws['B2'].font = Font(bold=True, size=18, color=INK)
    r = 4
    for t in lines:
        ws.cell(r, 2, t).font = Font(size=11, color='26323A')
        r += 1
    return ws, r + 1


# =================================================== PRÁCTICAS
def practicas(path):
    bk = Book()
    rng = np.random.default_rng(2026)
    scores = []
    for code, title, col, fn in TIPOS:
        for lvl in (1, 2, 3):
            sh = bk.sheet(f'{code}_N{lvl}', f'{title} · {NIV[lvl]}', col)
            fn(sh, lvl, np.random.default_rng(int(code[0]) * 10 + lvl))
            scores.append((sh.name, title, NIV[lvl], sh.finish(), sh.n_answers))
    ws, r = cover(bk, 'Prácticas autocorregidas', [
        'Cada tipo de ejercicio tiene tres niveles, con datos distintos en cada uno:',
        '   Nivel 1 · Con guía: la tabla de trabajo ya tiene las columnas y una fila de pistas.',
        '   Nivel 2 · Con cabeceras: tienes las columnas, pero sin pistas.',
        '   Nivel 3 · Sin ayuda: solo los datos y las preguntas, como en el examen.',
        'Escribe tus resultados en las celdas amarillas. Se ponen verdes si están bien y rojas si no;',
        'al lado aparece una pista cuando el error es uno de los típicos.',
        'Puedes hacer los cálculos donde quieras de la hoja (a la derecha hay sitio) y escribir solo el resultado,',
        'o poner directamente la fórmula en la celda amarilla. Los decimales con coma, como siempre.',
        'Las proporciones puedes darlas en decimal (0,25) o en porcentaje (25).'])
    hdr = ['Hoja', 'Ejercicio', 'Nivel', 'Aciertos']
    for j, h in enumerate(hdr):
        c = ws.cell(r, 2 + j if j < 3 else 5, h) if False else ws.cell(r, [2, 3, 4, 5][j], h)
        c.font = Font(bold=True, color=INK); c.fill = F_HEAD
    ws.column_dimensions['E'].width = 12
    for i, (nm, t, lv, ref, n) in enumerate(scores):
        rr = r + 1 + i
        c = ws.cell(rr, 2, nm); c.hyperlink = f"#'{nm}'!A1"; c.font = Font(color='0563C1', underline='single')
        ws.cell(rr, 3, t); ws.cell(rr, 4, lv.split('·')[1].strip())
        ws.cell(rr, 5, f'={ref}&" de {n}"').alignment = Alignment(horizontal='right')
    bk.save(path)
    return bk


# =================================================== SIMULACROS
def simulacros(path):
    bk = Book()
    sheets = []
    # ---- Simulacro 1 (el de la guía del tutor)
    sh = bk.sheet('S1_P1', 'Simulacro 1 · Pregunta 1 (2 puntos)', 't1', 'Días de baja en el último año de 25 trabajadores')
    t_discreta(sh, 3, None, data=[0, 2, 1, 0, 3, 1, 0, 0, 2, 5, 1, 0, 1, 2, 0, 9, 1, 0, 3, 1, 2, 0, 1, 4, 0],
               ctx={'var': 'Días de baja en el último año de 25 trabajadores', 'k': 2})
    sh.ask('¿Cuál es el único valor atípico? (límite superior C₃ + 1,5·IQR)', 9, 0, alts=[(5, 'Para ser atípico tiene que ser MAYOR que el límite (2 + 1,5·2 = 5). El 5 no lo es.')])
    sheets.append(sh)
    sh = bk.sheet('S1_P2', 'Simulacro 1 · Pregunta 2 (2 puntos)', 't1', 'Salario mensual de 110 empleados')
    t_intervalos(sh, 3, None, data=([900, 1200, 1500, 1800, 2400, 3600], [12, 30, 38, 24, 6]), ctx={'var': 'Salario mensual (€) de 110 empleados', 'hi': 10})
    sheets.append(sh)
    sh = bk.sheet('S1_P3', 'Simulacro 1 · Pregunta 3 (2 puntos)', 't2', 'Turno y satisfacción')
    t_conting(sh, 3, None, data=(['Mañana', 'Tarde', 'Noche'], ['Alta', 'Media', 'Baja'], [[18, 22, 10], [12, 20, 18], [6, 10, 24]], 'Turno', 'Satisfacción'))
    sheets.append(sh)
    sh = bk.sheet('S1_P4', 'Simulacro 1 · Pregunta 4 (2 puntos)', 't3', 'Formación y productividad')
    t_reg(sh, 3, None, data=([5, 8, 10, 12, 15, 18, 20, 25], [62, 65, 70, 71, 76, 80, 79, 88], 'Horas de formación (X)', 'Productividad en puntos (Y)', 22))
    sheets.append(sh)
    sh = bk.sheet('S1_P5', 'Simulacro 1 · Pregunta 5 (2 puntos)', 't4', 'Cesta de la compra y salario real')
    t_complejos(sh, 3, None, data=(['A', 'B', 'C'], [2.0, 5.0, 1.5], [50, 20, 80], [2.4, 5.5, 1.8], [45, 22, 70]))
    sheets.append(sh)
    # ---- Simulacro 2 (datos nuevos)
    rng = np.random.default_rng(77)
    sh = bk.sheet('S2_P1', 'Simulacro 2 · Pregunta 1 (2 puntos)', 't1', 'Hijos por empleado')
    t_discreta(sh, 3, None, data=[1, 0, 2, 2, 1, 3, 0, 1, 1, 2, 4, 0, 1, 2, 1, 0, 3, 1, 2, 1, 0, 2, 1, 6, 1, 0, 2, 1],
               ctx={'var': 'Número de hijos de 28 empleados', 'k': 1})
    sheets.append(sh)
    sh = bk.sheet('S2_P2', 'Simulacro 2 · Pregunta 2 (2 puntos)', 't1', 'Tiempo de desplazamiento al trabajo')
    t_intervalos(sh, 3, None, data=([0, 10, 20, 30, 45, 90], [14, 26, 31, 19, 10]), ctx={'var': 'Minutos de desplazamiento al trabajo de 100 personas', 'hi': 25})
    sheets.append(sh)
    sh = bk.sheet('S2_P3', 'Simulacro 2 · Pregunta 3 (2 puntos)', 't1', 'Salarios por sexo')
    t_subpob(sh, 3, None, data=(['Mujeres', 'Hombres'], [60, 90], [1980, 2240], [310, 420]), ctx={'unit': '€', 'ab': (1.02, 30), 'desc': 'suben un 2 % y además 30 €'})
    sheets.append(sh)
    sh = bk.sheet('S2_P4', 'Simulacro 2 · Pregunta 4 (2 puntos)', 't3', 'Absentismo y satisfacción')
    t_reg(sh, 3, None, data=([2, 3, 5, 6, 8, 9, 11, 12, 14], [8.4, 8.1, 7.2, 7.4, 6.1, 6.3, 5.0, 5.2, 4.1], 'Días de ausencia al año (X)', 'Satisfacción 0-10 (Y)', 10))
    sheets.append(sh)
    sh = bk.sheet('S2_P5', 'Simulacro 2 · Pregunta 5 (2 puntos)', 't4', 'Índices y tasas')
    t_indices(sh, 1, None)
    sheets.append(sh)
    refs = [(s.name, s.finish(), s.n_answers) for s in sheets]
    ws, r = cover(bk, 'Simulacros de examen autocorregidos', [
        'Dos exámenes de prueba de 5 preguntas (2 puntos cada una). Tiempo orientativo: 90 minutos cada uno.',
        'Haz cada simulacro de una sentada, con tus prácticas de clase y tu solucionario abiertos, como en el examen real.',
        'Escribe los resultados en las celdas amarillas: se corrigen solas. Haz los cálculos en la propia hoja.',
        'Las proporciones puedes darlas en decimal (0,25) o en porcentaje (25).'])
    for k, (tag, rng_) in enumerate([('Simulacro 1', refs[:5]), ('Simulacro 2', refs[5:])]):
        base = r + k * 9
        ws.cell(base, 2, tag).font = Font(bold=True, size=13, color=INK)
        ws.cell(base + 1, 2, 'Pregunta').font = Font(bold=True); ws.cell(base + 1, 3, 'Aciertos').font = Font(bold=True); ws.cell(base + 1, 4, 'Puntos').font = Font(bold=True)
        pts = []
        for i, (nm, ref, n) in enumerate(rng_):
            rr = base + 2 + i
            c = ws.cell(rr, 2, nm); c.hyperlink = f"#'{nm}'!A1"; c.font = Font(color='0563C1', underline='single')
            ws.cell(rr, 3, f'={ref}&" de {n}"')
            ws.cell(rr, 4, f'=ROUND({ref}/{n}*2,2)').number_format = '0.00'
            pts.append(f'D{rr}')
        ws.cell(base + 7, 3, 'Nota (sobre 10)').font = Font(bold=True)
        c = ws.cell(base + 7, 4, '=' + '+'.join(pts)); c.font = Font(bold=True, size=14, color='0F7A4E'); c.number_format = '0.0'
    bk.save(path)
    return bk


if __name__ == '__main__':
    import pickle
    b1 = practicas('/home/claude/xl2/Practicas_autocorregidas.xlsx')
    b2 = simulacros('/home/claude/xl2/Simulacros_autocorregidos.xlsx')
    pickle.dump({'p': b1.registry, 's': b2.registry}, open('/home/claude/xl2/registry.pkl', 'wb'))
    print('ok', len(b1.registry), len(b2.registry))
