import sys; sys.path.insert(0,'/home/claude/xl2'); sys.path.insert(0,'/home/claude/xl')
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
code1=open('/home/claude/xl/all.py').read()
code1=code1.replace("wb.save('/home/claude/xl/Estadistica_ejercicios_resueltos.xlsx');print('saved',len(wb.sheetnames))","")
ns1={'__name__':'res'};exec(code1,ns1)
wb=ns1['wb'];resolved=[(n,d) for n,d in ns1['sheets'] if n!='Simulacro']
wb.remove(wb['Índice']);wb.remove(wb['Simulacro'])
import re
def titles(f):
    return {int(m.group(1)):m.group(2) for m in re.finditer(r'^## Ejercicio (\d+)\. (.+)$',open(f).read(),re.M)}
TT={'B12':titles('/home/claude/build/alumna/03_boletin12.md')}
t34=open('/home/claude/build/alumna/04_tema3.md').read();TT['T3']={int(m.group(1)):m.group(2) for m in re.finditer(r'^## Ejercicio (\d+)\. (.+)$',t34,re.M)}
TT['T4']=titles('/home/claude/build/alumna/05_tema4.md')
def nice(n,d):
    g=n.split('_')[0];nums=[int(x) for x in re.findall(r'(?<=_)(?:Ej)?(\d\d)',n)]
    if g in TT and nums and n!='T3_Antig_Salario' and 'apuntes' not in n:
        return ' · '.join(f'Ej. {k}: {TT[g].get(k,"")}' for k in nums)+'  —  '+d
    return d
resolved=[(n,nice(n,d)) for n,d in resolved]

ns2={'WB':wb,'__name__':'plt'};exec(open('make_plantillas_lib.py').read(),ns2)
TOC=ns2['TOC'];INK='1F2A30'
F_HEAD=PatternFill('solid',fgColor='DDE5EA')
LINK=Font(color='0563C1',underline='single')
plant_names=[n for n,_,_ in TOC]
letter={n.split('_')[0]:n for n in plant_names}
# ---- Traductor: enlaces a plantilla y ejemplo resuelto
tr=wb['Traductor']
tr['E4']='Ejemplo resuelto';tr['E4'].font=Font(bold=True,color=INK);tr['E4'].fill=F_HEAD
tr.column_dimensions['E'].width=22
rules=[('Mediana',None),('moda','B12_Ej11'),('30 %','B12_Ej13'),('C₁ y C₃','T1_Cuartiles_apuntes'),('cuartil','B12_Ej12'),('Percentil','B12_Ej12'),
('más de"','B12_Ej06'),('"más de"','B12_Ej06'),('Coeficiente de variación','B12_Ej14'),('Cuasivarianza','T1_Hermanos'),('Varianza o desviación','B12_Ej08'),
('asimetría','T1_Hermanos'),('Curtosis','T1_Hermanos'),('Límites','T1_Hermanos'),('Tipificar','T1_Ponderada_Transf'),('Media ponderada','T1_Ponderada_Transf'),
('Varianza dentro','B12_Ej13'),('Varianza entre ÷','B12_Ej21'),('Sumar:','B12_Ej22'),('Reglas de cambio','B12_Ej22'),
('totales','T2_Contingencia'),('fila (o columna)','T2_Contingencia'),('fᵢⱼ con','B12_Ej18'),('Cada casilla','T2_Contingencia'),
('b = Sxy','T3_Ej01_02_04'),('Recta de Y','T3_Antig_Salario'),('pendiente','T3_Antig_Salario'),('Correlación','T3_Antig_Salario'),('R²','T3_Antig_Salario'),
('Sustituir','T3_Antig_Salario'),('Covarianza','T3_Antig_Salario'),
('Índice simple','T4_Ej04_05_06'),('anterior − 1','T4_Ej04_05_06'),('Dividir toda la serie','T4_Ej04_05_06'),('Multiplicar','T4_Ej01_02_03'),
('Media geométrica','T4_Ej07_08'),('(1 + anual)','T4_Ej07_08'),('A = (1','T4_Ej01_02_03'),('Laspeyres','T4_Ej09_Laspeyres'),('Paasche','T4_Ej09_Laspeyres'),
('√(Laspeyres','T4_Ej09_Laspeyres'),('IPC × 100','T4_Deflactar_apuntes'),('Comparar con','T4_Ej10_11_12'),('Media','T1_Hermanos')]
for r in range(5,tr.max_row+1):
    b=tr.cell(r,2).value
    if not b: continue
    tgt=None
    if str(b).startswith('Mediana'): tgt='T1_Hermanos'
    else:
        for k,v in rules:
            if v and k.lower() in str(b).lower(): tgt=v;break
    if tgt:
        c=tr.cell(r,5,tgt);c.hyperlink=f"#'{tgt}'!A1";c.font=LINK
    d=tr.cell(r,4)
    if d.value:
        first=str(d.value).split('/')[0].strip().split('_')[0].strip()
        if first in letter:
            d.hyperlink=f"#'{letter[first]}'!A1";d.font=LINK
# ---- enlace "volver" en todas las hojas
for ws in wb.worksheets:
    if ws.title=='Claves': continue
    col=10 if ws.title in plant_names else (6 if ws.title=='Traductor' else 9)
    c=ws.cell(1,col,'← Volver al inicio');c.hyperlink="#'Inicio'!A1";c.font=Font(color='0563C1',underline='single',bold=True)
# ---- Inicio
ini=wb.create_sheet('Inicio',0);ini.sheet_view.showGridLines=False
for k,v in {'A':3,'B':26,'C':62,'D':40}.items(): ini.column_dimensions[k].width=v
ini['B2']='Mi solucionario de Estadística';ini['B2'].font=Font(bold=True,size=20,color=INK)
ini['B3']='Ejercicios resueltos con fórmulas, calculadoras y traductor de enunciados, todo en un archivo.';ini['B3'].font=Font(italic=True,color='4F5D65')
steps=['Cómo usarlo en el examen:',
'1.  Lee la pregunta y busca una frase parecida en la hoja «Traductor» (Ctrl + B para buscar).',
'2.  Pulsa el enlace de la columna «Plantilla» (calculadora) o «Ejemplo resuelto».',
'3.  Copia la hoja antes de tocarla: clic derecho en la pestaña → Mover o copiar → marca «Crear una copia».',
'4.  En la copia, cambia los datos (celdas amarillas o azules) por los del examen. Todo se recalcula solo.',
'5.  Pasa el resultado a tu examen y escribe la frase de interpretación.',
'Si tus datos tienen más o menos filas que el ejemplo resuelto, usa mejor la plantilla: admite hasta 500 datos.']
for i,t in enumerate(steps):
    c=ini.cell(5+i,2,t);c.font=Font(bold=(i==0),size=11,color=INK if i==0 else '26323A')
r=14
def section(title,items,color):
    global r
    c=ini.cell(r,2,title);c.font=Font(bold=True,size=13,color='FFFFFF')
    for col in (2,3,4): ini.cell(r,col).fill=PatternFill('solid',fgColor=color)
    r+=1
    for nm,desc in items:
        c=ini.cell(r,2,nm);c.hyperlink=f"#'{nm}'!A1";c.font=LINK
        ini.cell(r,3,desc).alignment=Alignment(wrap_text=True);r+=1
    r+=1
section('Paso 1 · ¿Qué me están pidiendo?',[('Traductor','Frases típicas del enunciado → qué calcular, plantilla y ejemplo resuelto')],'1F2A30')
section('Calculadoras (pega tus datos)',[(n,t.split('·',1)[1].strip() if '·' in t else t) for n,t,d in TOC],'5B6B78')
grp={'T1':('Tema 1 · ejemplos de los apuntes','2F5FD0'),'T2':('Tema 2 · ejemplos de los apuntes','0F8A68'),'B12':('Boletín de los Temas 1 y 2','0F8A68'),'T3':('Tema 3 · regresión','C2185B'),'T4':('Tema 4 · números índices','A86F00')}
order=['T1','T2','B12','T3','T4']
for g in order:
    items=[(n,d) for n,d in resolved if n.split('_')[0]==g]
    if items: section('Resueltos · '+grp[g][0],items,grp[g][1])
# ---- orden de hojas
names=['Inicio','Traductor']+plant_names+[n for g in order for n,_ in resolved if n.split('_')[0]==g]
rest=[n for n in wb.sheetnames if n not in names]
wb._sheets=[wb[n] for n in names+rest]
for ws in wb.worksheets: ws.sheet_view.tabSelected=False
wb['Inicio'].sheet_view.tabSelected=True;wb.active=0
wb.save('/home/claude/xl2/Solucionario_estadistica.xlsx');print(len(wb.sheetnames),rest)
